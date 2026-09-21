"""Report templates and file formats — no network.

The same facts get written three times a week: a slide for the executives, a document for
the team, a sheet for whoever tracks the budget. The template is described in words and
the model turns that into a layout.

The split is the whole design and it is what most of these tests are about:

  **the model chooses the shape. The code writes every number.**

A model asked to "write the report" would restate figures in its own words, and a restated
figure is a new figure — close enough to look right and no longer traceable. So the spec
has nowhere to put content, and the composer reads every value out of the calculations it
was handed.

Run: pytest -q
"""

from __future__ import annotations

import datetime as dt
import json

from pmo_agent import render

TODAY = dt.date(2026, 9, 21)


def envelope(payload):
    return {
        "candidates": [{"content": {"parts": [{"text": json.dumps(payload, ensure_ascii=False)}]}}]
    }


def model(payload):
    def call(prompt):
        call.prompt = prompt
        return envelope(payload)

    return call


DATA = {
    "rollup": {
        "projects": [
            {
                "name": "Echo",
                "state": "risk",
                "headline": "2026-11-09",
                "headline_label": "achievable finish",
                "why": "the provider has not replied",
            },
            {
                "name": "Calm",
                "state": "ok",
                "headline": "2026-10-01",
                "headline_label": "achievable finish",
                "why": "",
            },
        ]
    },
    "schedule": {
        "risks": [
            {
                "key": "E-1",
                "summary": "Auth",
                "assignee": "Ada",
                "duedate": "2026-09-02",
                "predicted_finish": "2026-10-01",
                "delay_days": 21,
            },
        ]
    },
    "stalls": {
        "worst": [
            {
                "key": "E-1",
                "summary": "Auth",
                "window": {"days": 13},
                "explained_days": 3.5,
                "note": "3.5 of 13 accounted for",
            },
        ]
    },
    "questions": ["which task did 'authentication' mean?"],
}


class TestTheSpecHasNoContentInIt:
    def test_the_schema_offers_nowhere_to_put_findings(self):
        # Not a rule the model is asked to follow — there is no field for it
        fields = set(render.SPEC_SCHEMA["properties"])
        assert not fields & {"content", "rows", "text", "body", "findings", "summary_text"}

    def test_a_section_the_agent_cannot_fill_is_dropped(self):
        spec = render.clean({"sections": ["rollup", "profit_forecast", "risks"]})
        assert spec["sections"] == ["rollup", "risks"]
        assert spec["dropped"] == ["profit_forecast"]

    def test_the_dropped_ones_are_named_rather_than_silently_gone(self):
        # "this template has no risks section" and "the model asked for one I do not
        # have" are different facts and a reader needs to tell them apart
        spec = render.clean({"sections": ["invented"]})
        doc = render.compose(spec, DATA, TODAY)
        assert "invented" in doc["note"]

    def test_a_spec_with_no_usable_section_falls_back_rather_than_producing_nothing(self):
        assert render.clean({"sections": ["nope"]})["sections"] == list(render.DEFAULT)

    def test_the_row_limit_is_clamped_to_something_sane(self):
        assert render.clean({"rows_per_section": 9999})["rows_per_section"] == 50
        assert render.clean({"rows_per_section": 0})["rows_per_section"] == 5
        assert render.clean({"rows_per_section": "lots"})["rows_per_section"] == 5

    def test_what_was_asked_for_is_kept_next_to_what_was_understood(self):
        spec = render.template(
            "executive, one page", model({"sections": ["rollup"], "note": "a one page summary"})
        )
        assert spec["asked_for"] == "executive, one page"
        assert spec["understood"] == "a one page summary"


class TestWhenTheModelIsNoHelp:
    def test_an_empty_request_still_produces_a_template(self):
        assert render.template("")["sections"] == list(render.DEFAULT)

    def test_unreadable_json_falls_back_rather_than_failing(self):
        broken = {"candidates": [{"content": {"parts": [{"text": "{not json"}]}}]}
        assert render.parse_spec(broken)["sections"] == list(render.DEFAULT)

    def test_a_failed_call_says_so_and_still_produces_a_template(self):
        def boom(prompt):
            raise RuntimeError("the model call failed")

        spec = render.template("anything", boom)
        assert spec["sections"] == list(render.DEFAULT)
        assert "failed" in spec["error"]


class TestComposing:
    def spec(self, **kw):
        return render.clean({"sections": ["rollup", "risks", "stalls", "questions"], **kw})

    def test_every_value_comes_from_the_data_it_was_handed(self):
        doc = render.compose(self.spec(), DATA, TODAY)
        rollup = next(s for s in doc["sections"] if s["name"] == "rollup")
        assert rollup["rows"][0][0] == "Echo"
        assert rollup["rows"][0][2] == "2026-11-09"

    def test_only_trouble_leaves_out_what_is_on_track(self):
        doc = render.compose(self.spec(only_trouble=True), DATA, TODAY)
        rollup = next(s for s in doc["sections"] if s["name"] == "rollup")
        assert [r[0] for r in rollup["rows"]] == ["Echo"]

    def test_without_that_flag_everything_is_shown(self):
        doc = render.compose(self.spec(), DATA, TODAY)
        rollup = next(s for s in doc["sections"] if s["name"] == "rollup")
        assert [r[0] for r in rollup["rows"]] == ["Echo", "Calm"]

    def test_a_section_with_nothing_in_it_says_so_rather_than_disappearing(self):
        doc = render.compose(render.clean({"sections": ["signals"]}), DATA, TODAY)
        assert doc["sections"][0]["rows"] == []
        assert "signals" in doc["note"]

    def test_rows_beyond_the_limit_are_counted_not_dropped_silently(self):
        many = {
            "rollup": {
                "projects": [{"name": f"P{i}", "state": "risk", "headline": "x"} for i in range(9)]
            }
        }
        doc = render.compose(
            render.clean({"sections": ["rollup"], "rows_per_section": 3}), many, TODAY
        )
        assert len(doc["sections"][0]["rows"]) == 3
        assert doc["sections"][0]["more"] == 6

    def test_the_unaccounted_part_survives_into_the_report(self):
        # The whole reason this project exists must not be lost at the last step
        doc = render.compose(render.clean({"sections": ["stalls"]}), DATA, TODAY)
        assert "3.5 of 13" in doc["sections"][0]["rows"][0][3]

    def test_the_sections_keep_the_order_the_spec_gave(self):
        spec = render.clean({"sections": ["questions", "rollup"]})
        doc = render.compose(spec, DATA, TODAY)
        assert [s["name"] for s in doc["sections"]] == ["questions", "rollup"]


class TestMarkdown:
    def test_it_writes_a_table_per_section(self):
        text = render.to_markdown(
            render.compose(render.clean({"sections": ["risks"]}), DATA, TODAY)
        )
        assert "| task | owner | due | forecast | slip |" in text
        assert "E-1 Auth" in text

    def test_a_pipe_in_a_value_does_not_break_the_table(self):
        data = {"questions": ["is it a | or a slash?"]}
        text = render.to_markdown(
            render.compose(render.clean({"sections": ["questions"]}), data, TODAY)
        )
        assert r"a \| or" in text

    def test_the_date_and_the_audience_are_on_it(self):
        spec = render.clean({"sections": ["rollup"], "audience": "executives"})
        text = render.to_markdown(render.compose(spec, DATA, TODAY))
        assert "executives" in text and "2026-09-21" in text


class TestTheFileFormats:
    def compose(self):
        return render.compose(
            render.clean({"sections": ["rollup", "risks", "questions"], "title": "Weekly"}),
            DATA,
            TODAY,
        )

    def test_markdown_comes_back_as_text_not_a_file(self):
        got = render.write(self.compose(), "md")
        assert got["format"] == "md" and got["text"].startswith("# Weekly")

    def test_a_word_document_is_written(self, tmp_path):
        got = render.write(self.compose(), "docx", str(tmp_path / "r.docx"))
        assert got["bytes"] > 0

        from docx import Document

        read_back = Document(got["path"])
        text = "\n".join(p.text for p in read_back.paragraphs)
        assert "Weekly" in text
        assert any("E-1 Auth" in c.text for t in read_back.tables for r in t.rows for c in r.cells)

    def test_a_spreadsheet_is_written_with_a_sheet_per_section(self, tmp_path):
        got = render.write(self.compose(), "xlsx", str(tmp_path / "r.xlsx"))
        from openpyxl import load_workbook

        book = load_workbook(got["path"])
        assert {"rollup", "risks", "questions", "about"} <= set(book.sheetnames)
        assert book["rollup"].cell(2, 1).value == "Echo"

    def test_the_spreadsheet_records_what_was_asked_for(self, tmp_path):
        got = render.write(self.compose(), "xlsx", str(tmp_path / "r.xlsx"))
        from openpyxl import load_workbook

        rows = list(load_workbook(got["path"])["about"].values)
        assert any(r[0] == "note" for r in rows)

    def test_a_deck_is_written_with_a_slide_per_section(self, tmp_path):
        got = render.write(self.compose(), "pptx", str(tmp_path / "r.pptx"))
        from pptx import Presentation

        deck = Presentation(got["path"])
        assert len(deck.slides) == 4, "a cover plus one per section"

    def test_an_unknown_format_is_refused_by_name(self):
        got = render.write(self.compose(), "pdf")
        assert "unknown format: pdf" in got["error"]

    def test_the_format_is_read_leniently(self):
        assert render.write(self.compose(), ".MD")["format"] == "md"

    def test_the_default_filename_carries_the_date_and_the_title(self, tmp_path):
        doc = self.compose()
        name = render._default_path(doc, "docx")
        assert name.endswith(".docx") and "2026-09-21" in name and "Weekly" in name


class TestTheModelNeverWritesANumber:
    def test_the_composer_reads_nothing_but_the_data_it_is_given(self):
        # Composing against empty data must produce empty tables, never filler
        spec = render.clean({"sections": list(render.SECTIONS)})
        doc = render.compose(spec, {}, TODAY)
        assert all(s["rows"] == [] for s in doc["sections"])

    def test_the_spec_is_kept_with_the_report_so_it_can_be_checked(self):
        doc = render.compose(render.clean({"sections": ["rollup"]}), DATA, TODAY)
        assert doc["spec"]["sections"] == ["rollup"]
