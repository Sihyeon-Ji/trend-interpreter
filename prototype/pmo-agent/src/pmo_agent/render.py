"""Reports in the shape somebody asked for, in the format they need to hand over.

The same facts get written three times a week: a slide for the executives, a document for
the team, a sheet for whoever tracks the budget. Same numbers, three shapes, and all three
are typed by hand.

So a template is described in words — *"executive, one page, only the projects in trouble,
three lines each"* — and the model turns that into a **spec**: which sections, in what
order, how long, filtered how. Then the code fills it.

That split is the whole design, and it is the same one the rest of this codebase uses:

  **the model chooses the shape. The code writes every number.**

A model asked to "write the report" would restate figures in its own words, and a restated
figure is a new figure — close enough to look right and no longer traceable to the
calculation that produced it. So the spec has no content in it. It cannot: the schema has
nowhere to put any.

Formats are `md`, `docx`, `xlsx` and `pptx`. Each is a different arrangement of the same
composed rows, never a different set of facts.
"""

from __future__ import annotations

import datetime as _dt
import json
import os

from pmo_agent import llm

# What a template is allowed to ask for. A section the agent cannot fill from a real
# calculation is not on this list, which is why the model cannot invent one.
SECTIONS = {
    "rollup": "every project on one line, worst first",
    "risks": "tasks forecast to miss their date",
    "stalls": "why work stopped, with the evidence",
    "issues": "the open issue register",
    "signals": "where the board disagrees with the code",
    "capacity": "who the project is waiting on",
    "queue": "an operations project's backlog and what recurs",
    "questions": "what the agent could not settle and needs a person for",
}
FORMATS = ("md", "docx", "xlsx", "pptx")
DEFAULT = ("rollup", "risks", "stalls", "questions")

SPEC_SYSTEM = """You turn a request for a report into a layout. Rules:
1) Choose only from the section names you are given. Never invent one.
2) You are choosing shape, not content: which sections, their order, how many rows each.
   Do not write any findings, numbers or sentences. There is nowhere to put them.
3) If the request names an audience (executives, the team, finance), set audience to it.
4) only_trouble means: leave out anything that is on track.
5) If the request is vague, prefer fewer sections and shorter rows."""

SPEC_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "audience": {"type": "string"},
        "sections": {"type": "array", "items": {"type": "string"}},
        "rows_per_section": {"type": "integer"},
        "only_trouble": {"type": "boolean"},
        "note": {"type": "string", "description": "what you understood the request to ask for"},
    },
    "required": ["sections"],
}


def call(prompt: str) -> dict:
    """The one model call, injectable, so every test below runs with no network."""
    listed = "\n".join(f"  {name}: {what}" for name, what in SECTIONS.items())
    return llm.generate(
        {
            "contents": [
                {
                    "role": "user",
                    "parts": [{"text": f"Sections available:\n{listed}\n\nRequest:\n{prompt}"}],
                }
            ],
            "systemInstruction": {"parts": [{"text": SPEC_SYSTEM}]},
            "generationConfig": {
                "temperature": 0,
                "responseMimeType": "application/json",
                "responseSchema": SPEC_SCHEMA,
            },
        }
    )


def parse_spec(resp: dict, prompt: str = "") -> dict:
    """Model response → a spec the composer can trust. Anything unknown is dropped.

    Dropping rather than failing: a template that lost a section still produces a report,
    and a section nobody can fill is worse than one that is missing.
    """
    raw = "".join(
        p.get("text", "")
        for p in (((resp.get("candidates") or [{}])[0].get("content") or {}).get("parts") or [])
    ).strip()
    try:
        data = json.loads(raw) if raw else {}
    except json.JSONDecodeError:
        data = {}
    return clean(data, prompt)


def clean(data: dict, prompt: str = "") -> dict:
    """A spec with only things that exist in it."""
    wanted = [s for s in (data.get("sections") or []) if s in SECTIONS]
    dropped = [s for s in (data.get("sections") or []) if s not in SECTIONS]
    try:
        asked = int(data.get("rows_per_section"))  # type: ignore[arg-type]
    except (TypeError, ValueError):
        asked = 0
    # 0 and "lots" both mean "did not say", which is 5. Only a real number is clamped.
    rows = max(1, min(asked, 50)) if asked > 0 else 5
    return {
        "title": str(data.get("title") or "Project report").strip()[:120],
        "audience": str(data.get("audience") or "").strip()[:60],
        "sections": wanted or list(DEFAULT),
        "rows_per_section": rows,
        "only_trouble": bool(data.get("only_trouble")),
        "asked_for": prompt.strip()[:300],
        "understood": str(data.get("note") or "").strip()[:200],
        # Named so a reader can tell "this template has no risks section" from
        # "the model asked for one I do not have".
        "dropped": dropped,
    }


def template(prompt: str, caller=call) -> dict:
    """A request in words → a layout. No content, by construction."""
    if not (prompt or "").strip():
        return clean({}, "")
    try:
        return parse_spec(caller(prompt), prompt)
    except Exception as e:
        out = clean({}, prompt)
        out["error"] = str(e)[:160]
        return out


# --------------------------------------------------------------------- composing
def _trouble(row: dict) -> bool:
    return row.get("state") in ("risk", "watch", "unknown")


def compose(spec: dict, data: dict, today: _dt.date | None = None) -> dict:
    """The spec plus the calculations → titled tables. Every value comes from `data`.

    Nothing here writes a sentence of its own. If a number is on the page, some other
    module computed it and this only chose where it sits.
    """
    today = today or _dt.date.today()
    limit = spec.get("rows_per_section", 5)
    only = spec.get("only_trouble")
    out = []

    # A spec that came from somebody's own document carries the original order, with the
    # sections nothing can fill left in place. Otherwise it is just the list of sections.
    walk = spec.get("order") or [{"heading": "", "section": n} for n in spec.get("sections") or []]

    for step in walk:
        name = step.get("section") or ""
        if not name:
            # Their heading, kept where it was and marked. Dropping it lets a reader
            # believe their template was applied when it was not, and filling it is the
            # worst thing a report can do.
            why = next(
                (
                    u.get("why", "")
                    for u in spec.get("unfillable") or []
                    if u.get("heading") == step.get("heading")
                ),
                "nothing here can fill this",
            )
            out.append(
                {
                    "name": "",
                    "title": step.get("heading", ""),
                    "head": [],
                    "rows": [],
                    "more": 0,
                    "unfillable": why,
                }
            )
            continue
        rows, head = [], []
        if name == "rollup":
            head = ["project", "state", "headline", "what it is", "why"]
            source = (data.get("rollup") or {}).get("projects") or []
            rows = [
                [
                    r["name"],
                    r["state"],
                    r["headline"],
                    r.get("headline_label", ""),
                    r.get("why", ""),
                ]
                for r in source
                if not only or _trouble(r)
            ]
        elif name == "risks":
            head = ["task", "owner", "due", "forecast", "slip"]
            rows = [
                [
                    f"{r['key']} {r['summary']}",
                    r.get("assignee", ""),
                    r.get("duedate", ""),
                    r.get("predicted_finish", ""),
                    f"{r.get('delay_days', 0)}d",
                ]
                for r in (data.get("schedule") or {}).get("risks") or []
            ]
        elif name == "stalls":
            head = ["task", "silent", "accounted for", "why"]
            rows = [
                [
                    f"{r['key']} {r['summary']}",
                    f"{r['window']['days']}d",
                    f"{r.get('explained_days', 0)} of {r['window']['days']}",
                    r.get("note", ""),
                ]
                for r in (data.get("stalls") or {}).get("worst") or []
            ]
        elif name == "issues":
            head = ["issue", "task", "owner", "days", "since"]
            rows = [
                [
                    r.get("detail", ""),
                    r.get("key", ""),
                    r.get("owner", ""),
                    r.get("days", 0),
                    r.get("first_seen", ""),
                ]
                for r in (data.get("issues") or {}).get("open") or []
            ]
        elif name == "signals":
            head = ["task", "kind", "what happened"]
            rows = [
                [f"{m['key']} {m['summary']}", m.get("kind", ""), m.get("detail", "")]
                for m in (data.get("signals") or {}).get("mismatches") or []
            ]
        elif name == "capacity":
            head = ["owner", "days left", "used", "free from"]
            rows = [
                [r["assignee"], r["busy_days"], f"{round(r['utilization'] * 100)}%", r["free_at"]]
                for r in (data.get("capacity") or {}).get("resources") or []
            ]
        elif name == "queue":
            head = ["fault", "times filed", "still open"]
            rows = [
                [r["fault"], r["times"], r["open"]]
                for r in (data.get("health") or {}).get("recurring") or []
            ]
        elif name == "questions":
            head = ["the agent could not settle this"]
            rows = [[q] for q in data.get("questions") or []]

        out.append(
            {
                "name": name,
                # Their heading when the template came from their own document, ours
                # otherwise. A report that arrives under different headings gets
                # reformatted by hand, which is the tedium this was meant to remove.
                "title": (spec.get("headings") or {}).get(name) or SECTIONS[name],
                "head": head,
                "rows": rows[:limit],
                "more": max(0, len(rows) - limit),
            }
        )

    return {
        "title": spec.get("title") or "Project report",
        "audience": spec.get("audience", ""),
        "today": today.isoformat(),
        "preamble": spec.get("preamble") or [],
        "sections": out,
        "spec": spec,
        "note": _note(out, spec),
    }


def _note(sections: list, spec: dict) -> str:
    """One line about what this report is and is not."""
    empty = [s["name"] for s in sections if not s["rows"] and s["name"]]
    cannot = [s["title"] for s in sections if s.get("unfillable")]
    said = f"{len(sections)} sections"
    if spec.get("only_trouble"):
        said += ", only what is in trouble"
    if empty:
        said += f". Nothing to report under: {', '.join(empty)}"
    if spec.get("dropped"):
        said += f". Asked for sections this agent does not have: {', '.join(spec['dropped'])}"
    if cannot:
        said += (
            f". {len(cannot)} section(s) from the source template are left empty because "
            f"nothing here can fill them: {', '.join(cannot[:3])}"
        )
    return said + "."


# --------------------------------------------------------------------- writing
def to_markdown(doc: dict) -> str:
    lines = [f"# {doc['title']}", ""]
    if doc.get("audience"):
        lines.append(f"*For {doc['audience']} · {doc['today']}*")
    else:
        lines.append(f"*{doc['today']}*")
    lines.append("")
    for section in doc["sections"]:
        lines += [f"## {section['title']}", ""]
        if section.get("unfillable"):
            lines += [f"*\u2014 {section['unfillable']}. This agent did not fill it.*", ""]
            continue
        if not section["rows"]:
            lines += ["Nothing to report.", ""]
            continue
        lines.append("| " + " | ".join(section["head"]) + " |")
        lines.append("|" + "|".join("---" for _ in section["head"]) + "|")
        for row in section["rows"]:
            lines.append("| " + " | ".join(str(c).replace("|", "\\|") for c in row) + " |")
        if section["more"]:
            lines.append("")
            lines.append(f"*and {section['more']} more*")
        lines.append("")
    lines += ["---", f"*{doc['note']}*"]
    return "\n".join(lines)


def to_docx(doc: dict, path: str) -> str:
    from docx import Document

    out = Document()
    out.add_heading(doc["title"], level=0)
    who = f"For {doc['audience']} · " if doc.get("audience") else ""
    out.add_paragraph(f"{who}{doc['today']}")
    for section in doc["sections"]:
        out.add_heading(section["title"], level=1)
        if section.get("unfillable"):
            out.add_paragraph(f"\u2014 {section['unfillable']}. This agent did not fill it.")
            continue
        if not section["rows"]:
            out.add_paragraph("Nothing to report.")
            continue
        table = out.add_table(rows=1, cols=len(section["head"]))
        table.style = "Light Grid Accent 1"
        for cell, name in zip(table.rows[0].cells, section["head"], strict=False):
            cell.text = str(name)
        for row in section["rows"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=False):
                cell.text = str(value)
        if section["more"]:
            out.add_paragraph(f"and {section['more']} more")
    out.add_paragraph(doc["note"])
    out.save(path)
    return path


def to_xlsx(doc: dict, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    for section in doc["sections"]:
        sheet = book.create_sheet(section["name"][:31])
        sheet.append(section["head"] or ["(nothing)"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in section["rows"]:
            sheet.append([str(c) for c in row])
        for column in range(1, max(1, len(section["head"])) + 1):
            longest = max(
                (len(str(sheet.cell(r, column).value or "")) for r in range(1, sheet.max_row + 1)),
                default=10,
            )
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = min(
                60, max(12, longest + 2)
            )
    about = book.create_sheet("about")
    about.append(["title", doc["title"]])
    about.append(["date", doc["today"]])
    about.append(["note", doc["note"]])
    about.append(["asked for", doc["spec"].get("asked_for", "")])
    book.save(path)
    return path


def to_pptx(doc: dict, path: str) -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = doc["title"]
    who = f"For {doc['audience']}\n" if doc.get("audience") else ""
    cover.placeholders[1].text = f"{who}{doc['today']}"

    for section in doc["sections"]:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = section["title"]
        if section.get("unfillable"):
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(8.8), Inches(1))
            box.text_frame.text = f"\u2014 {section['unfillable']}. This agent did not fill it."
            continue
        if not section["rows"]:
            box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(8.8), Inches(1))
            box.text_frame.text = "Nothing to report."
            continue
        # One row per line rather than a table: a slide that has to be squinted at is
        # worse than one that says less.
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(8.8), Inches(4.8))
        frame = box.text_frame
        frame.word_wrap = True
        for index, row in enumerate(section["rows"][:8]):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.text = " · ".join(str(c) for c in row if str(c))[:180]
            para.font.size = Pt(13)
        if section["more"]:
            para = frame.add_paragraph()
            para.text = f"and {section['more']} more"
            para.font.size = Pt(11)
    deck.save(path)
    return path


WRITERS = {"docx": to_docx, "xlsx": to_xlsx, "pptx": to_pptx}


def write(doc: dict, fmt: str = "md", path: str = "") -> dict:
    """Write the composed report. `md` comes back as text; the rest as a file."""
    fmt = (fmt or "md").strip().lower().lstrip(".")
    if fmt not in FORMATS:
        return {"error": f"unknown format: {fmt} (one of {', '.join(FORMATS)})"}
    if fmt == "md":
        return {"format": "md", "text": to_markdown(doc)}
    target = path or _default_path(doc, fmt)
    os.makedirs(os.path.dirname(target) or ".", exist_ok=True)
    try:
        WRITERS[fmt](doc, target)
    except ImportError as e:
        return {"error": f"{fmt} needs a library that is not installed: {e}"}
    return {"format": fmt, "path": target, "bytes": os.path.getsize(target)}


def _default_path(doc: dict, fmt: str) -> str:
    from pmo_agent.env import reports_dir

    safe = "".join(c if c.isalnum() or c in "-_" else "-" for c in doc["title"]).strip("-")[:48]
    return str(reports_dir() / f"{doc['today']}-{safe or 'report'}.{fmt}")
